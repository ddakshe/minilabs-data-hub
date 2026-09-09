"""K-apt 엑셀 3종 → aptcost/ 산출물.

앱이 답하는 질문은 하나다: **"우리 집 관리비, 많이 내는 편인가?"**
그 답은 백분위인데, 백분위를 내려면 비교군 전체가 있어야 한다. 그래서 앱에서
API 를 부르는 구조가 성립하지 않는다 — 사용자 한 명이 수백 콜을 쓴다.
여기서 미리 계산해 시군구 단위로 쪼개 넣고, 앱은 자기 시군구 파일 하나만 읽는다.

비교를 공정하게 만드는 규칙 세 가지:

  1. **원/㎡ 로 환산한다.** 세대당 총액은 평수 차이로 왜곡된다.
     분모는 `관리비부과면적`(면적정보 엑셀). 이게 없으면 그 단지는 버린다.
  2. **12개월 평균을 쓴다.** 관리비는 계절이 지배한다(난방비는 6월 37만 → 7월 9천).
     특정 달을 고정하면 그 달 미제출 단지가 통째로 빠지는 문제도 같이 해결된다.
  3. **공용관리비와 개별사용료를 나눈다.** 공용은 관리사무소 몫이고 개별은 우리 집
     사용량이다. 합쳐서 한 줄로 주면 "난방 많이 썼다"와 "관리비가 비싸다"가 섞인다.

산출물:
    aptcost/meta.json              기준일·연도·단지 수·월별 커버리지
    aptcost/complexes.json         단지 마스터 (앱의 검색용)
    aptcost/stats/{code}.json      시군구별 지표 + 백분위 경계
"""

from __future__ import annotations

import argparse
import collections
import json
import statistics
import warnings
from datetime import datetime, timezone, timedelta
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

KST = timezone(timedelta(hours=9))

# 관리비 엑셀에서 쓸 컬럼. 52개 전부 넣으면 파일이 커지고, 앱이 실제로 보여주는 건
# 이 아홉이다. 세부 항목(인건비·피복비…)은 "왜 비싼가"를 파고들 때 추가한다.
#
# 🚨 **컬럼 이름을 여러 개 받는다.** K-apt 가 서식을 예고 없이 바꾼다 —
#    2026-09-07 추출본의 `공용관리비계`·`일반관리비계` 가 2026-09-09 추출본에서
#    `공용관리비(합계)`·`일반관리비(합계)` 로 바뀌었고 `수선유지비(합계)` 가 신설됐다
#    (51열 → 52열). 이틀 만이다.
#    이름을 하나만 두고 `.get()` 에 기본값을 주면 **에러 없이 전부 0** 이 되어
#    모든 단지가 "해당 없음" 으로 나온다. 그래서 후보를 나열하고, 하나도 못 찾으면
#    아래에서 **죽는다**. 조용히 틀린 값을 내보내느니 멈추는 게 낫다.
COST_FIELDS: dict[str, tuple[str, ...]] = {
    "common": ("공용관리비(합계)", "공용관리비계"),
    "guard": ("경비비",),
    "clean": ("청소비",),
    "elev": ("승강기유지비",),
    "repair": ("수선비",),
    "indiv": ("개별사용료계", "개별사용료(합계)"),
    "heat": ("난방비(전용)",),
    "elec": ("전기료(전용)",),
    "ltrf": ("장충금 월부과액",),
}


def _pick(H: dict, names: tuple[str, ...]) -> int | None:
    """후보 이름 중 실제로 있는 컬럼의 인덱스."""
    for n in names:
        if n in H:
            return H[n]
    return None


def _num(v) -> float:
    """엑셀이 '2.9551852E7' 같은 문자열로 준다."""
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _rows(path: Path):
    """(헤더, 데이터 이터레이터). 0행은 안내문이라 건너뛴다.

    read_only 모드에서 dimension 이 1x1 로 잘못 온다 → reset_dimensions 필수.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()
    it = ws.iter_rows(values_only=True)
    next(it)  # 안내문
    header = next(it)
    return list(header), it


def load_area(path: Path) -> dict[str, dict]:
    """단지코드 → 관리비부과면적·세대수·평형 목록.

    🚨 이 엑셀은 **단지당 여러 행**이다 — 주거전용면적 구간(59㎡·84㎡…)마다 한 행씩.
       `관리비부과면적` 은 단지 총계라 어느 행에서 읽어도 같지만, `세대수` 는
       **그 면적 구간의 세대수**다. 첫 행만 쓰면 단지 세대수가 아니라 한 평형의
       세대수를 쓰게 된다(그랑빌 3,000세대급이 725로 나왔다).
       세대수는 구간을 전부 더한다.

    평형 목록(`sz`)을 같이 싣는다. 앱이 "우리 집 평형" 을 고르게 하려면 필요하다 —
    자유 입력을 받으면 오타와 검증이 따라오는데, 그 단지에 실제로 있는 평형만
    보여주면 탭 한 번으로 끝난다.

    ⚠️ 지표의 분모는 **관리비부과면적**이고 사용자가 아는 건 **전용면적**이다.
       둘은 다르다(부과면적이 공용분을 포함해 더 크다). 단지 총계 비율
       `관리비부과면적 / 주거전용면적합` 을 `far` 로 실어 앱이 환산하게 한다.
    """
    header, it = _rows(path)
    H = {h: i for i, h in enumerate(header)}
    out: dict[str, dict] = {}
    for r in it:
        code = r[H["단지코드"]]
        if not code:
            continue
        area = _num(r[H["관리비부과면적"]])
        hh = int(_num(r[H["세대수"]]))
        priv = _num(r[H["주거전용면적(세부)"]])
        cur = out.get(code)
        if cur is None:
            if area <= 0:
                continue
            total_priv = _num(r[H["주거전용면적(단지합계)"]])
            cur = out[code] = {
                "area": round(area, 1),
                "hh": hh,
                "far": round(area / total_priv, 3) if total_priv > 0 else None,
                "sz": [],
            }
        else:
            cur["hh"] += hh
        if priv > 0 and hh > 0:
            cur["sz"].append([round(priv, 1), hh])
    for v in out.values():
        # 같은 평형이 여러 행으로 쪼개져 오기도 한다 — 합치고 면적순으로 세운다
        merged: dict[float, int] = {}
        for a, n in v["sz"]:
            merged[a] = merged.get(a, 0) + n
        v["sz"] = [[a, n] for a, n in sorted(merged.items())]
    return out


def load_basis(path: Path) -> dict[str, dict]:
    """단지코드 → 비교군을 나누는 속성들.

    난방방식이 특히 중요하다. 개별난방 단지는 난방비가 관리비에 안 잡혀 0 으로
    보이는데, 그건 싼 게 아니라 다른 데로 나가는 것이다. 같은 난방방식끼리
    비교하지 않으면 앱이 틀린 답을 자신 있게 준다.
    """
    header, it = _rows(path)
    H = {h: i for i, h in enumerate(header)}

    def g(r, name):
        i = H.get(name)
        return r[i] if i is not None else None

    out: dict[str, dict] = {}
    for r in it:
        code = g(r, "단지코드")
        if not code or code in out:
            continue
        used = str(g(r, "사용승인일") or "")
        out[code] = {
            "_name": g(r, "단지명"),
            "_sido": g(r, "시도"),
            "_sgg": g(r, "시군구"),
            "_dong": g(r, "동리") or g(r, "읍면") or "",
            "heatType": g(r, "난방방식"),
            "mgmtType": g(r, "관리방식"),
            "builtYear": int(used[:4]) if used[:4].isdigit() else None,
            "guardCnt": int(_num(g(r, "경비관리-인원"))),
            "cleanCnt": int(_num(g(r, "청소관리-인원"))),
            # 금액만 보여주면 "비싸다" 가 비난이 된다. 경비를 몇 명 두는지, 승강기가
            # 몇 대인지를 같이 주면 **선택의 결과**로 읽힌다 — 경비 인원이 많은 건
            # 낭비가 아니라 서비스 수준일 수 있다. 이 뉘앙스가 없으면 앱이
            # 관리사무소 분쟁 도구가 된다.
            "elevCnt": sum(
                int(_num(g(r, c)))
                for c in (
                    "승강기(승객용)",
                    "승강기(화물용)",
                    "승강기(승객+화물)",
                    "승강기(장애인)",
                    "승강기(비상용)",
                    "승강기(기타)",
                )
            ),
            "parkCnt": int(_num(g(r, "총주차대수"))),
            "cctvCnt": int(_num(g(r, "CCTV대수"))),
            "kind": g(r, "단지분류"),
            # 세대수의 진실원. 면적정보의 구간 합계와 어긋나면 이쪽을 쓴다.
            "hh": int(_num(g(r, "세대수"))),
        }
    return out


def load_cost(path: Path) -> tuple[dict[str, dict], collections.Counter, dict[str, dict]]:
    """단지코드 → 항목별 12개월 합계. 월별 커버리지도 같이 센다."""
    header, it = _rows(path)
    H = {h: i for i, h in enumerate(header)}
    idx = {k: _pick(H, names) for k, names in COST_FIELDS.items()}
    missing = {k: COST_FIELDS[k] for k, i in idx.items() if i is None}
    if missing:
        raise RuntimeError(
            f"관리비 엑셀에서 컬럼을 못 찾았다: {missing}\n"
            f"실제 헤더: {[h for h in header if h]}\n"
            "→ K-apt 가 서식을 바꿨을 수 있다. COST_FIELDS 에 새 이름을 추가한다."
        )
    common_i = _pick(H, COST_FIELDS["common"])
    indiv_i = _pick(H, COST_FIELDS["indiv"])

    agg: dict[str, dict] = {}
    months = collections.Counter()
    where: dict[str, dict] = {}
    for r in it:
        code = r[H["단지코드"]]
        if not code:
            continue
        ym = str(r[H["발생년월(YYYYMM)"]])
        months[ym] += 1
        a = agg.get(code)
        if a is None:
            a = agg[code] = {"n": 0, "by_month": {}, **{k: 0.0 for k in COST_FIELDS}}
            where[code] = {
                "name": r[H["단지명"]],
                "sido": r[H["시도"]],
                "sgg": r[H["시군구"]],
                "dong": r[H["동리"]] or r[H["읍면"]] or "",
            }
        a["n"] += 1
        # 월별 공용·개별. 12개월 평균만 실으면 겨울 트리거("난방비가 두 배인데
        # 우리만 이런가")에 답할 수 없다. 계절이 이 도메인의 절반이다.
        a["by_month"][ym[4:]] = (_num(r[common_i]), _num(r[indiv_i]))
        for key, i in idx.items():
            a[key] += _num(r[i])
    return agg, months, where


def load_region_index(path: Path) -> dict[tuple[str, str], str]:
    """(시도, 시군구) → 시군구 5자리 코드.

    `realestate/region-master.json` 을 그대로 쓰되 표기 차이를 흡수한다.
    두 시스템이 일반구를 다르게 적는다 —
        region-master : "성남시 분당구"  (법정동코드 표기)
        K-apt 엑셀    : "성남분당구"
    이걸 안 맞추면 분당·동탄·기흥 같은 아파트 밀집지가 통째로 빠진다(실측 4,700 단지).
    세종은 시군구가 없어 엑셀에서 빈 값으로 온다.
    """
    regions = json.loads(path.read_text(encoding="utf-8"))["regions"]
    idx: dict[tuple[str, str], str] = {}
    for r in regions:
        sido, sgg, code = r["sido"], r["sigungu"], r["code"]
        idx.setdefault((sido, sgg), code)
        idx.setdefault((sido, sgg.replace("시 ", "")), code)  # 성남시 분당구 → 성남분당구
        if sgg == sido:  # 세종특별자치시
            idx.setdefault((sido, ""), code)
            idx.setdefault((sido, None), code)
    return idx


def percentiles(values: list[float]) -> dict:
    """앱이 '상위 몇 %'를 그리는 데 필요한 경계값."""
    s = sorted(values)
    if not s:
        return {}

    def at(p: float) -> float:
        return round(s[min(len(s) - 1, int(len(s) * p))], 1)

    return {
        "n": len(s),
        "p10": at(0.10), "p25": at(0.25), "p50": at(0.50),
        "p75": at(0.75), "p90": at(0.90),
        "mean": round(statistics.fmean(s), 1),
    }


# 지역 마스터는 저장소 루트 기준으로 찾는다. --out 의 부모에서 유추하면
# --out 을 저장소 밖으로 주는 순간 조용히 깨진다(리허설에서 실측).
REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_REGIONS = REPO_ROOT / "realestate" / "region-master.json"


def recover_trend(out: Path) -> tuple[dict[str, dict[int, list]], int | None]:
    """직전 산출물에서 추이를 되살린다.

    🚨 이게 없으면 **정기 실행이 과거를 지운다.** 월 1회 실행은 올해와 작년 엑셀만
       받는데(과거 연도를 매달 500MB 씩 받을 이유가 없다) 빌더는 raw 폴더에 있는
       연도로만 추이를 만든다. 병합하지 않으면 2018~2023 이 매달 날아간다.
       (리허설에서 산출물이 14MB → 11MB 로 줄어드는 것으로 실측)

    직전 기준 연도의 대표 지표도 추이로 접어 넣는다. 기준 연도가 2025→2026 으로
    넘어가면 2025 는 더 이상 raw 에 없지만 추이에는 남아야 한다.
    """
    stats_dir = out / "stats"
    if not stats_dir.is_dir():
        return {}, None
    recovered: dict[str, dict[int, list]] = collections.defaultdict(dict)
    prev_base: int | None = None
    for f in stats_dir.glob("*.json"):
        try:
            doc = json.loads(f.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        prev_base = doc.get("year") or prev_base
        for code, rec in doc.get("complexes", {}).items():
            for y, v in (rec.get("trend") or {}).items():
                recovered[code][int(y)] = v
            if prev_base is not None and "common" in rec:
                recovered[code].setdefault(
                    int(doc["year"]), [rec["common"], rec["indiv"], rec.get("months", 0)]
                )
    return recovered, prev_base


def build(
    raw: Path,
    out: Path,
    year: int,
    trend_years: list[int] | None = None,
    regions_path: Path | None = None,
) -> None:
    cost_files = sorted(raw.glob(f"*관리비정보_{year}.xlsx"))
    basis_files = sorted(raw.glob("*기본정보.xlsx"))
    area_files = sorted(raw.glob("*면적정보.xlsx"))
    for label, files in (("관리비", cost_files), ("기본정보", basis_files), ("면적정보", area_files)):
        if not files:
            raise SystemExit(f"{label} 엑셀이 {raw} 에 없다 — kapt_files.py 로 먼저 받는다")

    area = load_area(area_files[-1])
    basis = load_basis(basis_files[-1])
    agg, months, where = load_cost(cost_files[-1])

    # 추이. 기준 연도는 전체 지표를 싣지만 나머지 연도는 공용·개별 두 값만 얹는다.
    # 9개 연도를 전부 같은 모양으로 넣으면 시군구 파일이 9배가 된다.
    #
    # ⚠️ 연도끼리는 **같은 단지끼리만** 비교해야 한다. 공개 의무 대상이
    #    2024-10-25 부터 100세대 이상으로 넓어져서, 전국 평균을 연도별로 늘어놓으면
    #    값이 변한 게 아니라 단지 구성이 바뀐 것을 보게 된다.
    #    단지별 추이는 이 문제가 없다 — 그래서 추이를 단지 단위로만 만든다.
    # ⚠️ 진행 중인 연도는 개월 수가 모자란다. months 를 같이 넣어 앱이 판단하게 한다.
    trend, prev_base = recover_trend(out)
    trend = collections.defaultdict(dict, trend)
    if trend:
        kept = sorted({y for d in trend.values() for y in d})
        print(f"  이전 산출물에서 추이 복구: {len(trend):,} 단지 · 연도 {kept}"
              + (f" (직전 기준 {prev_base})" if prev_base else ""))
    trend_meta: dict[int, dict] = {}
    for ty in sorted(trend_years or []):
        if ty == year:
            continue
        files = sorted(raw.glob(f"*관리비정보_{ty}.xlsx"))
        if not files:
            print(f"  추이: {ty} 엑셀 없음 — 건너뜀")
            continue
        t_agg, t_months, _ = load_cost(files[-1])
        n_hit = 0
        for code, a in t_agg.items():
            if code not in area or a["n"] == 0:
                continue
            m2 = area[code]["area"]
            trend[code][ty] = [
                round(a["common"] / a["n"] / m2, 1),
                round(a["indiv"] / a["n"] / m2, 1),
                a["n"],
            ]
            n_hit += 1
        trend_meta[ty] = {"complexes": n_hit, "months": len(t_months)}
        print(f"  추이: {ty} · 단지 {n_hit:,} · {len(t_months)}개월")

    code_of = load_region_index(regions_path or DEFAULT_REGIONS)

    complexes: dict[str, dict] = {}
    by_region: dict[str, list[str]] = collections.defaultdict(list)
    unmatched = collections.Counter()

    for code, a in agg.items():
        if code not in area or a["n"] == 0:
            continue
        m2 = area[code]["area"]
        w = where[code]
        region = code_of.get((w["sido"], w["sgg"]))
        if region is None:
            unmatched[(w["sido"], w["sgg"])] += 1
            continue
        rec = {
            "name": w["name"],
            "sido": w["sido"],
            "sgg": w["sgg"],
            "dong": w["dong"],
            "region": region,
            # 기본정보 세대수를 우선한다(면적 구간 합계는 누락 구간이 있을 수 있다)
            "hh": int(basis.get(code, {}).get("hh") or area[code]["hh"]),
            "area": m2,
            "months": a["n"],
            **{k: round(a[k] / a["n"] / m2, 1) for k in COST_FIELDS},
            **{k: v for k, v in basis.get(code, {}).items() if v is not None and k != "hh" and not k.startswith("_")},
        }
        # 평형 목록과 부과/전용 면적비 — 앱이 "우리 집 평형" 을 고르는 데 쓴다
        if area[code].get("sz"):
            rec["sz"] = area[code]["sz"]
        if area[code].get("far"):
            rec["far"] = area[code]["far"]

        # 월별 원/㎡ — "01".."12" 키. 단지당 숫자 24개라 파일이 크게 늘지 않는다.
        by_month = a.get("by_month") or {}
        if by_month:
            rec["m"] = {
                mm: [round(cg / m2, 1), round(ig / m2, 1)]
                for mm, (cg, ig) in sorted(by_month.items())
            }
        if code in trend:
            # [공용, 개별, 개월수] — 연도를 키로. 앱은 개월수로 진행 중인 해를 가린다.
            # 기준 연도는 대표 지표로 따로 실리므로 추이에서 뺀다(복구분에 섞여 들어온다).
            rec["trend"] = {
                str(y): v for y, v in sorted(trend[code].items()) if y != year
            }
        complexes[code] = rec
        by_region[region].append(code)

    # 🚨 관리비 자료가 없는 단지도 목록에는 띄운다.
    #
    #    K-apt 에 단지는 등록돼 있는데 관리비를 안 올린 곳이 있다(공개 의무가 없는
    #    소규모 단지, 최근 가입 단지). 그런 단지를 목록에서 통째로 빼면 사용자는
    #    "우리 아파트가 왜 없지" 하고 **검색이 고장난 줄 안다.**
    #    이름은 띄우고 자료가 없다는 사실을 말해주는 게 정직하고, 헤매지 않게 한다.
    no_data = 0
    for code, b in basis.items():
        if code in complexes:
            continue
        region = code_of.get((b.get("_sido"), b.get("_sgg")))
        if region is None or not b.get("_name"):
            continue
        complexes[code] = {
            "name": b["_name"],
            "sido": b["_sido"],
            "sgg": b["_sgg"],
            "dong": b["_dong"],
            "region": region,
            "hh": int(b.get("hh") or 0),
            "nd": 1,  # no data — 앱이 이걸 보고 회색 처리한다
        }
        by_region[region].append(code)
        no_data += 1

    # 시군구별 파일. 앱은 자기 시군구 하나만 받는다.
    stats_dir = out / "stats"
    stats_dir.mkdir(parents=True, exist_ok=True)
    for old in stats_dir.glob("*.json"):
        old.unlink()
    for region, codes in by_region.items():
        rows = {c: complexes[c] for c in codes}
        # 자료 없는 단지는 분포에 넣지 않는다 — 0 으로 세면 중앙값이 끌려 내려간다
        withData = [r for r in rows.values() if not r.get("nd")]
        dist = {
            key: percentiles([r[key] for r in withData if r.get(key, 0) > 0])
            for key in COST_FIELDS
        }
        (stats_dir / f"{region}.json").write_text(
            json.dumps({"region": region, "year": year, "dist": dist, "complexes": rows},
                       ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )

    # 앱이 부팅하며 받는 시도→시군구 목록. realestate/region-master.json 을 앱이 직접
    # 읽게 하면 디렉터리 간 결합이 생기고, 실거래 쪽 사정으로 목록이 바뀌면 관리비 앱이
    # 같이 흔들린다. 여기서 **실제로 데이터가 있는 시군구만** 따로 굽는다(253개 · 15KB).
    # 표기는 K-apt 엑셀 것을 그대로 쓴다(성남분당구·전남광주통합특별시). region-master 의
    # 표기로 갈아끼우면 매핑이 한 겹 더 생기고, 화면에 뜨는 이름과 데이터가 어긋난다.
    regions_out = []
    for region, codes in sorted(by_region.items()):
        sample = complexes[codes[0]]
        regions_out.append(
            {"code": region, "sido": sample["sido"], "sigungu": sample["sgg"], "n": len(codes)}
        )
    (out / "regions.json").write_text(
        json.dumps(regions_out, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )

    # 앱의 단지 검색용 — 지표는 빼고 이름·주소만.
    (out / "complexes.json").write_text(
        json.dumps(
            {c: {"n": r["name"], "r": r["region"], "d": r["dong"], "h": r["hh"]}
             for c, r in complexes.items()},
            ensure_ascii=False, separators=(",", ":"),
        ),
        encoding="utf-8",
    )

    haveData = [r for r in complexes.values() if not r.get("nd")]

    trend_count: collections.Counter = collections.Counter()
    trend_months: dict[int, int] = {}
    for r in haveData:
        for y, v in (r.get("trend") or {}).items():
            trend_count[int(y)] += 1
            trend_months[int(y)] = max(trend_months.get(int(y), 0), v[2])

    nationwide = {
        key: percentiles([r[key] for r in haveData if r.get(key, 0) > 0])
        for key in COST_FIELDS
    }
    (out / "meta.json").write_text(
        json.dumps(
            {
                "fetchedAt": datetime.now(KST).isoformat(timespec="seconds"),
                "source": "K-apt 자료실 (국토교통부·한국부동산원)",
                "sourceFile": cost_files[-1].name,
                "year": year,
                "complexes": len(complexes),
                "regions": len(by_region),
                "monthlyRows": {m: months[m] for m in sorted(months)},
                # 복구분까지 포함한 최종 추이 현황. 새로 읽은 연도만 세면
                # 정기 실행 메타에 과거가 사라진 것처럼 보인다.
                "trendYears": {
                    str(y): {"complexes": n, "months": trend_months[y]}
                    for y, n in sorted(trend_count.items())
                },
                "trendNote": (
                    "연도 비교는 같은 단지끼리만 한다. 공개 의무가 2024-10-25 부터 "
                    "100세대 이상으로 넓어져 전국 평균의 연도 비교는 단지 구성 변화를 본다."
                ),
                "nationwide": nationwide,
                "unmatchedRegions": {f"{s} {g}": n for (s, g), n in unmatched.most_common()},
            },
            ensure_ascii=False, indent=2,
        ),
        encoding="utf-8",
    )

    print(f"단지 {len(complexes):,} (자료 있음 {len(complexes) - no_data:,} · 자료 없음 {no_data:,})"
          f" · 시군구 {len(by_region)} · {stats_dir}")
    if unmatched:
        print("지역 매칭 실패:", dict(unmatched.most_common(8)))
    print("전국 공용관리비 원/㎡·월:", nationwide["common"])


def main() -> None:
    ap = argparse.ArgumentParser(description="K-apt 엑셀 → aptcost/")
    ap.add_argument("--raw", default="_raw/aptcost", help="엑셀이 있는 디렉터리")
    ap.add_argument("--out", default="aptcost", help="산출 디렉터리")
    ap.add_argument(
        "--year",
        default="auto",
        help="기준 연도 (전체 지표를 싣는 해). 'auto' 는 12개월이 다 찬 가장 최근 연도",
    )
    ap.add_argument(
        "--regions",
        default="",
        help=f"시군구 마스터 경로 (기본 {DEFAULT_REGIONS.relative_to(REPO_ROOT)})",
    )
    ap.add_argument(
        "--trend",
        default="",
        help="추이로 얹을 연도. 'all' 또는 '2018,2019,…'. 기준 연도는 자동 제외",
    )
    args = ap.parse_args()
    raw = Path(args.raw)
    have = sorted(int(p.stem.rsplit("_", 1)[-1]) for p in raw.glob("*관리비정보_*.xlsx"))
    years = have if args.trend == "all" else [int(y) for y in args.trend.split(",") if y.strip()]

    if args.year == "auto":
        # 진행 중인 연도를 기준으로 삼으면 계절이 반쪽이 된다. 관리비는 난방비 때문에
        # 상반기만 평균 내면 겨울 쪽으로 쏠린다. 12개월이 다 찬 최근 연도를 쓴다.
        base = None
        for y in reversed(have):
            _, months, _ = load_cost(sorted(raw.glob(f"*관리비정보_{y}.xlsx"))[-1])
            if len(months) >= 12:
                base = y
                break
        if base is None:
            raise SystemExit("12개월이 다 찬 연도가 없다 — --year 로 직접 지정한다")
        print(f"기준 연도(auto): {base}")
    else:
        base = int(args.year)

    build(raw, Path(args.out), base, years, Path(args.regions) if args.regions else None)


if __name__ == "__main__":
    main()
